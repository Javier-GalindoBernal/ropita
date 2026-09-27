"""Servidor de Ropita: exporta ropita_db.xlsx a JSON, sirve el frontend (docs/) y conecta con Ollama.

Uso:  python3 server.py      →  http://localhost:8000
Variables opcionales: OLLAMA_URL (default http://localhost:11434), OLLAMA_MODEL (default llama3.2), PORT.
"""
import json
import os
import urllib.error
import urllib.request
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.path.join(BASE_DIR, "ropita_db.xlsx")
DOCS = os.path.join(BASE_DIR, "docs")
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.2")
PORT = int(os.environ.get("PORT", 8000))


def leer_hoja(wb, nombre):
    filas = list(wb[nombre].iter_rows(values_only=True))
    return [dict(zip(filas[0], f)) for f in filas[1:] if any(f)]


def exportar_json():
    """Convierte el Excel en docs/ropita_data.json, que es lo que lee el frontend (también en GitHub Pages)."""
    wb = load_workbook(EXCEL, read_only=True)
    datos = {"outfits": leer_hoja(wb, "Outfits"),
             "preguntas": sorted(leer_hoja(wb, "Preguntas"), key=lambda p: p["orden"])}
    with open(os.path.join(DOCS, "ropita_data.json"), "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Base exportada: {len(datos['outfits'])} outfits, {len(datos['preguntas'])} preguntas")


SISTEMA = """Eres Ropita, un asistente de moda amable que habla español.
Recomiendas outfits según el clima, la ocasión, el momento del día, el día de la semana, la edad y el género.
Basa tus respuestas en el PERFIL y el OUTFIT BASE que vienen de la base de datos; puedes complementarlos con
alternativas, colores y consejos. Responde en máximo 120 palabras, con viñetas cuando ayuden.
Habla directo a la persona: nunca menciones las palabras "PERFIL", "OUTFIT BASE" ni "base de datos".
Si la persona es menor de edad, mantén recomendaciones apropiadas para su edad.
Si te preguntan algo que no sea moda o vestuario, redirige amablemente la conversación."""


def chat_ollama(perfil, outfit, mensaje, historial):
    contexto = f"PERFIL: {json.dumps(perfil, ensure_ascii=False)}\nOUTFIT BASE: {json.dumps(outfit, ensure_ascii=False)}"
    mensajes = [{"role": "system", "content": SISTEMA + "\n\n" + contexto}]
    mensajes += historial[-8:]
    mensajes.append({"role": "user", "content": mensaje})
    cuerpo = json.dumps({"model": OLLAMA_MODEL, "messages": mensajes, "stream": False}).encode()
    req = urllib.request.Request(f"{OLLAMA_URL}/api/chat", data=cuerpo, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.load(r)["message"]["content"]


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=DOCS, **kw)

    def responder(self, datos, status=200):
        cuerpo = json.dumps(datos, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(cuerpo)))
        self.end_headers()
        self.wfile.write(cuerpo)

    def do_POST(self):
        datos = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or "{}")
        respuestas = datos.get("respuestas", {})

        if self.path == "/api/chat":
            try:
                texto = chat_ollama(respuestas, datos.get("outfit"), datos.get("mensaje", ""), datos.get("historial", []))
                return self.responder({"respuesta": texto})
            except (urllib.error.URLError, TimeoutError, KeyError) as e:
                return self.responder({"error": f"No pude conectar con Ollama ({OLLAMA_URL}, modelo {OLLAMA_MODEL}): {e}"}, 503)

        self.responder({"error": "Ruta no encontrada"}, 404)


if __name__ == "__main__":
    exportar_json()
    print(f"Ropita listo en http://localhost:{PORT}  (Ollama: {OLLAMA_URL}, modelo {OLLAMA_MODEL})")
    ThreadingHTTPServer(("", PORT), Handler).serve_forever()
